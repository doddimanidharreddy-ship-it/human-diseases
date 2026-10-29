"""
Generates the professional 13-slide PowerPoint presentation matching
Vel Tech Minor Project Review 1 specifications for:
DISEASE PREDICTION SYSTEM USING DECISION TREES
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = pptx.Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # completely blank
    assets_dir = os.path.join(os.path.dirname(__file__), "assets")
    logo_path = os.path.join(assets_dir, "veltech_logo.png")
    arch_img_path = os.path.join(assets_dir, "system_architecture_diagram.png")

    def add_footer(slide, slide_num):
        # Footer text frame
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"30 September 2026       BATCH NO: MI1301       DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING       {slide_num}"
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(120, 120, 120)

    # ==================== SLIDE 1: TITLE SLIDE ====================
    slide1 = prs.slides.add_slide(blank_layout)
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(5.2), Inches(0.4), width=Inches(3.0))

    tx1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(2.2))
    tf1 = tx1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nSCHOOL OF COMPUTING\n10214CS601- MINOR PROJECT – I\nACADEMIC YEAR – 2026-2027 (SUMMER SEMESTER)"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.bold = True
    p1.font.size = Pt(20)
    p1.font.color.rgb = RGBColor(15, 30, 60)

    # Presented By
    tx_pres = slide1.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(6.0), Inches(2.0))
    tf_pres = tx_pres.text_frame
    p_pres_head = tf_pres.paragraphs[0]
    p_pres_head.text = "PRESENTED BY"
    p_pres_head.font.bold = True
    p_pres_head.font.size = Pt(16)
    p_pres_head.font.color.rgb = RGBColor(20, 40, 80)

    students = [
        "1. Bugga Gangadhar Kali Prasad (VTU 28564)(24UECS0049)",
        "2. Allala Bhavana (VTU 27999)(24UECS0008)",
        "3. Godi Mouni Pranathi (VTU 28298)(24UECS0107)"
    ]
    for s in students:
        p = tf_pres.add_paragraph()
        p.text = s
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(40, 40, 40)

    # Supervised By
    tx_sup = slide1.shapes.add_textbox(Inches(7.5), Inches(4.5), Inches(5.0), Inches(2.0))
    tf_sup = tx_sup.text_frame
    p_sup_head = tf_sup.paragraphs[0]
    p_sup_head.text = "SUPERVISED BY"
    p_sup_head.font.bold = True
    p_sup_head.font.size = Pt(16)
    p_sup_head.font.color.rgb = RGBColor(20, 40, 80)

    p_sup = tf_sup.add_paragraph()
    p_sup.text = "Dr. J. Raja (Associate Professor)"
    p_sup.font.size = Pt(15)
    p_sup.font.bold = True
    p_sup.font.color.rgb = RGBColor(40, 40, 40)

    add_footer(slide1, 1)

    # Helper for Content Slides
    def add_content_slide(title_text, subtitle_text, bullet_points, slide_num):
        slide = prs.slides.add_slide(blank_layout)
        
        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.8))
        tf = t_box.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(10, 25, 55)

        curr_top = Inches(1.3)
        if subtitle_text:
            s_box = slide.shapes.add_textbox(Inches(0.8), curr_top, Inches(11.733), Inches(0.5))
            sf = s_box.text_frame
            sp = sf.paragraphs[0]
            sp.text = subtitle_text
            sp.font.size = Pt(20)
            sp.font.bold = True
            sp.font.color.rgb = RGBColor(20, 40, 80)
            curr_top = Inches(1.9)

        # Bullets
        c_box = slide.shapes.add_textbox(Inches(0.8), curr_top, Inches(11.733), Inches(5.0))
        cf = c_box.text_frame
        cf.word_wrap = True

        for idx, item in enumerate(bullet_points):
            if idx == 0:
                p = cf.paragraphs[0]
            else:
                p = cf.add_paragraph()
            
            p.text = "•  " + item
            p.font.size = Pt(16)
            p.font.color.rgb = RGBColor(40, 40, 40)
            p.space_after = Pt(12)

        add_footer(slide, slide_num)
        return slide

    # ==================== SLIDE 2: AGENDA ====================
    agenda_bullets = [
        "Introduction",
        "Literature Survey",
        "Problem Statement",
        "Proposed Methodology",
        "Proposed System Architecture",
        "Implementation Plan/Timeline"
    ]
    add_content_slide("DISEASE PREDICTION SYSTEM USING DECISION TREES", "AGENDA", agenda_bullets, 2)

    # ==================== SLIDE 3: INTRODUCTION - PROBLEM CONTEXT ====================
    intro_p1_bullets = [
        "There is no simple, accessible and transparent tool that lets an ordinary user check their risk for multiple diseases from basic health parameters in one place.",
        "Most existing student projects and tools handle only a single disease, are not deployed as usable applications, or are hard to extend.",
        "Many advanced systems use black-box models (such as deep learning) that give predictions without explaining why, which reduces trust in a medical setting.",
        "Symptom-checker apps are mostly closed-source, rule-based and symptom-only, and rarely use lab values for disease-specific risk.",
        "Many existing prediction tools give only a 'positive' or 'negative' result with no explanation, so users and health workers cannot judge whether to trust it."
    ]
    add_content_slide("Introduction", "Problem Statement Context:", intro_p1_bullets, 3)

    # ==================== SLIDE 4: INTRODUCTION - PROJECT GOAL ====================
    intro_p2_bullets = [
        "Project Goal: To develop an intelligent, application-based system that estimates the risk of major chronic and acute diseases and helps users and health workers make timely health decisions using Decision Tree machine learning.",
        "The system aims to:",
        "Predict the risk of 9 disease domains: Diabetes, Heart Disease, Kidney Disease, Liver Disease, Breast Cancer, Stroke, Parkinson's, Hypertension, and Infectious Syndromes.",
        "Suggest likely differential conditions from 18 user-selected symptoms.",
        "Explain each prediction through a transparent, verifiable Decision Tree decision path.",
        "Provide timely recommendations, including risk level (Low to Critical) and when to consult a doctor.",
        "Reduce delays in early detection and manual report interpretation."
    ]
    add_content_slide("Introduction", None, intro_p2_bullets, 4)

    # ==================== SLIDE 5: INTRODUCTION - OVERVIEW ====================
    intro_p3_bullets = [
        "Disease Prediction System Using Decision Trees is an explainable machine learning clinical decision support platform covering 9 specialized disease domains based on patient vitals, lab biomarkers, and physical symptoms.",
        "The system deploys optimized Decision Tree classification algorithms trained on clinical benchmark datasets. When clinical parameters are submitted, the model analyzes the input features and follows a sequence of decision rules to predict disease likelihood.",
        "By combining multi-disease machine learning, explainable decision paths, and a deployable web application, the proposed system helps reduce delays in early detection, lessen the manual effort of interpreting health reports, and support preventive healthcare."
    ]
    add_content_slide("Introduction", "Brief Overview & Scope:", intro_p3_bullets, 5)

    # ==================== SLIDE 6: LITERATURE SURVEY ====================
    lit_bullets = [
        "Traditional diagnosis: Doctors use patient history, symptoms, physical examination, laboratory tests and medical reports to identify diseases.",
        "Rule-based systems: Some healthcare applications use predefined medical rules to identify possible diseases from symptoms and patient information.",
        "Machine learning systems: Existing research applies algorithms such as Decision Trees, Random Forest, SVM, KNN and Logistic Regression for disease prediction.",
        "Single-disease prediction: Many systems focus on only one disease, such as diabetes, heart disease, kidney disease or liver disease.",
        "Symptom-based prediction: Some web-based systems let users enter symptoms and receive a possible disease prediction from a trained model."
    ]
    add_content_slide("Literature Survey", "Existing Systems and Related Works:", lit_bullets, 6)

    # ==================== SLIDE 7: PROBLEM STATEMENT ====================
    ps_bullets = [
        "Limited Disease Coverage: Many existing systems are designed to predict only one specific disease, such as diabetes or heart disease.",
        "Dependence on Dataset Quality: Prediction accuracy depends heavily on the quality, size, and completeness of the medical dataset used for training.",
        "Limited Patient Information: Some systems use only a small number of symptoms or medical parameters, which may not provide enough information for reliable prediction.",
        "Lack of Interpretability: Complex machine learning models may provide predictions without clearly explaining how the result was obtained.",
        "No Unified Prediction Platform: Separate systems are often required for different diseases, making the process less convenient for users.",
        "Data Preprocessing Challenges: Missing, inconsistent, or imbalanced medical data can negatively affect model performance."
    ]
    add_content_slide("Problem Statement", "Limitations of Existing Solutions:", ps_bullets, 7)

    # ==================== SLIDE 8: PROPOSED METHODOLOGY ====================
    meth_bullets = [
        "Multi-disease prediction: Predicts Diabetes, Heart, Kidney, Liver, Breast Cancer, Stroke, Parkinson's, Hypertension, and 16 acute symptom conditions in one unified portal.",
        "Data collection and preprocessing: Takes health parameters (vitals, lab panels, acoustic biomarkers, symptoms) and cleans data by handling missing values, encoding, and range normalization.",
        "Decision tree classification: Trains on medical data with regularized depth to create understandable clinical rules and predict disease.",
        "Interpretable prediction: Shows clear root-to-leaf decision structures, unlike complex black-box neural networks.",
        "Model evaluation: Assessed using accuracy, precision, recall, F1-score and confusion matrix across all models.",
        "User-friendly interface and awareness: Lets users enter data without technical knowledge and shows results with advice to consult a doctor, not as a final diagnosis."
    ]
    add_content_slide("Proposed Methodology", None, meth_bullets, 8)

    # ==================== SLIDE 9: PROPOSED SYSTEM ARCHITECTURE ====================
    slide9 = prs.slides.add_slide(blank_layout)
    t9_box = slide9.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.8))
    tf9 = t9_box.text_frame
    p9 = tf9.paragraphs[0]
    p9.text = "Proposed System Architecture"
    p9.font.size = Pt(28)
    p9.font.bold = True
    p9.font.color.rgb = RGBColor(10, 25, 55)

    if os.path.exists(arch_img_path):
        slide9.shapes.add_picture(arch_img_path, Inches(2.2), Inches(1.3), width=Inches(8.9))

    add_footer(slide9, 9)

    # ==================== SLIDE 10: SUMMARY OF MODULE - 1 ====================
    summary_bullets = [
        "The Disease Prediction System Using Decision Trees provides comprehensive screening across 9 specialized clinical domains.",
        "It directly resolves the challenge of delayed diagnosis and fragmented single-disease tools.",
        "The system analyzes physiological parameters (vitals, blood panels, tumor morphology, voice biomarkers, and symptoms) to estimate risk.",
        "Decision Tree models ensure full algorithmic transparency, risk level estimation (Low to Critical), and step-by-step decision paths.",
        "The proposed system provides actionable lifestyle recommendations and encourages early medical consultation."
    ]
    add_content_slide("Summary of Module -1", None, summary_bullets, 10)

    # ==================== SLIDE 11: SOFTWARE & HARDWARE REQUIREMENTS ====================
    req_bullets = [
        "Operating System: Windows 10/11",
        "Programming Language: Python",
        "Development Environment: Visual Studio Code / Jupyter Notebook",
        "AI/ML: Scikit-learn (Decision Tree Classifier)",
        "Libraries: NumPy, Pandas, Matplotlib, Seaborn, Joblib",
        "Database: SQLite / MySQL (for user history, optional)",
        "User Interface: Streamlit / Flask",
        "Browser: Google Chrome / Microsoft Edge"
    ]
    add_content_slide("Software & Hardware Requirements Specification", None, req_bullets, 11)

    # ==================== SLIDE 12: IMPLEMENTATION PLAN / TIMELINE ====================
    timeline_bullets = [
        "Week 1-2: Problem identification, literature survey, and analysis of disease and dataset requirements in prediction.",
        "Week 3-4: System design, data preprocessing, and development of the first disease models (heart disease and diabetes) using Decision Trees.",
        "Week 5-6: Development of the liver and kidney disease models and the symptom-based module, with evaluation using accuracy, precision, recall and F1-score.",
        "Week 7-8: Integration into a single web application with risk level and recommendations, system testing, and final implementation."
    ]
    add_content_slide("Implementation Plan/Timeline", "Key Milestones:", timeline_bullets, 12)

    # ==================== SLIDE 13: THANK YOU ====================
    slide13 = prs.slides.add_slide(blank_layout)
    tx13 = slide13.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.333), Inches(2.5))
    tf13 = tx13.text_frame
    p13 = tf13.paragraphs[0]
    p13.text = "THANK YOU"
    p13.alignment = PP_ALIGN.CENTER
    p13.font.bold = True
    p13.font.size = Pt(44)
    p13.font.color.rgb = RGBColor(10, 25, 55)

    p13_sub = tf13.add_paragraph()
    p13_sub.text = "\nQuestions & Discussions\nDepartment of Computer Science & Engineering\nVel Tech R&D Institute of Science and Technology"
    p13_sub.alignment = PP_ALIGN.CENTER
    p13_sub.font.size = Pt(18)
    p13_sub.font.color.rgb = RGBColor(90, 90, 90)

    add_footer(slide13, 13)

    # Save to both project dir and Downloads
    out_path_project = os.path.join(os.path.dirname(__file__), "DISEASE_PREDICTION_SYSTEM_REVIEW_1.pptx")
    out_path_downloads = r"C:\Users\manid\Downloads\CSE-MINOR_PROJECT-1_DISEASE_PREDICTION.pptx"

    prs.save(out_path_project)
    prs.save(out_path_downloads)

    print(f"[OK] Saved presentation to: {out_path_project}")
    print(f"[OK] Saved presentation to: {out_path_downloads}")

if __name__ == "__main__":
    create_presentation()
